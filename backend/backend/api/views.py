from django.shortcuts import render
from rest_framework import viewsets, filters, status
from rest_framework.permissions import IsAuthenticated, AllowAny
from rest_framework.decorators import action
from rest_framework.permissions import AllowAny
from rest_framework.response import Response
from django_filters.rest_framework import DjangoFilterBackend
from .models import *
from .serializers import *
from .permissions import IsOwnerOrAdmin
from django.contrib.auth.models import Group
from rest_framework.generics import ListAPIView
from rest_framework.views import APIView
from rest_framework.parsers import MultiPartParser, FormParser, JSONParser
from django.db.models import Count
from decimal import ROUND_HALF_UP, Decimal
from django.utils.timezone import now
from django.utils import timezone
from django.core.exceptions import ValidationError as DjangoValidationError
import pandas as pd
import random
import string
import requests
from django.core.cache import cache

# Generer fichier ABP128
from openpyxl import load_workbook
from io import BytesIO
from django.http import FileResponse

from .serializers import _get_deadline, validate_password_custom


# ============================================================
# BASE VIEWSET — hérité par tous les autres ViewSets
# ============================================================

class BaseViewSet(viewsets.ModelViewSet):
    permission_classes = [IsAuthenticated, IsOwnerOrAdmin]
    filter_backends = [DjangoFilterBackend, filters.SearchFilter, filters.OrderingFilter]

    def perform_create(self, serializer):
        serializer.save(
            user_created=self.request.user,
            user_updated=self.request.user
        )

    def perform_update(self, serializer):
        serializer.save(user_updated=self.request.user)

    @action(detail=False, methods=['delete'])
    def bulk_delete(self, request):
        ids = request.data.get('ids', [])

        if not isinstance(ids, list) or not ids:
            return Response(
                {"error": "Veuillez fournir une liste d'IDs valide et non vide"},
                status=status.HTTP_400_BAD_REQUEST
            )

        if not all(isinstance(i, int) for i in ids):
            return Response(
                {"error": "Tous les IDs doivent être des entiers"},
                status=status.HTTP_400_BAD_REQUEST
            )

        queryset = self.get_queryset().filter(id__in=ids)
        found_count = queryset.count()

        if found_count == 0:
            return Response(
                {"error": "Aucun élément trouvé avec les IDs fournis"},
                status=status.HTTP_404_NOT_FOUND
            )

        deleted_count, _ = queryset.delete()

        return Response({
            "message": f"{deleted_count} élément(s) supprimé(s) avec succès",
            "deleted_count": deleted_count
        }, status=status.HTTP_200_OK)


# ============================================================
# USERS & GROUPES
# ============================================================

class UserPosteViewSet(BaseViewSet):
    queryset = UserPoste.objects.select_related(
        'user', 'poste', 'statut'
    ).order_by('id')
    serializer_class = UserPosteSerializer
    filterset_fields = ['user', 'poste', 'statut']  # ✅ déjà présent

    def perform_update(self, serializer):
        """✅ Vérifie que le record appartient au même user avant de sauvegarder."""
        instance = serializer.instance
        new_user = serializer.validated_data.get('user', instance.user)
        if instance.user_id != new_user.id:
            from rest_framework.exceptions import PermissionDenied
            raise PermissionDenied(
                "Impossible de changer le user d'un UserPoste existant."
            )
        super().perform_update(serializer)


class UserCategoriesalarieViewSet(BaseViewSet):
    queryset = UserCategoriesalarie.objects.select_related(
        'user', 'categoriesalarie', 'statut'
    ).order_by('id')
    serializer_class = UserCategoriesalarieSerializer
    filterset_fields = ['user', 'categoriesalarie', 'statut']  # ✅ déjà présent

    def perform_update(self, serializer):
        instance = serializer.instance
        new_user = serializer.validated_data.get('user', instance.user)
        if instance.user_id != new_user.id:
            from rest_framework.exceptions import PermissionDenied
            raise PermissionDenied(
                "Impossible de changer le user d'un UserCategoriesalarie existant."
            )
        super().perform_update(serializer)


class UserAllergieViewSet(BaseViewSet):
    queryset = UserAllergie.objects.select_related('user').order_by('id')
    serializer_class = UserAllergieSerializer
    filterset_fields = ['user']  # ✅ déjà présent

    def perform_update(self, serializer):
        instance = serializer.instance
        new_user = serializer.validated_data.get('user', instance.user)
        if instance.user_id != new_user.id:
            from rest_framework.exceptions import PermissionDenied
            raise PermissionDenied(
                "Impossible de changer le user d'un UserAllergie existant."
            )
        super().perform_update(serializer)


class UserViewSet(BaseViewSet):
    # queryset factice requis par le router DRF pour déterminer le basename.
    # La vraie logique est dans get_queryset() ci-dessous.
    queryset = User.objects.none()
    serializer_class = UserSerializer
    filterset_fields = ['nom', 'prenom', 'email', 'statut', 'is_active']
    search_fields    = ['nom', 'prenom', 'email', 'username']
    ordering_fields  = ['nom', 'created_at']

    def get_queryset(self):
        from django.db.models import Prefetch
        return User.objects.select_related(
            'statut'
        ).prefetch_related(
            Prefetch('userservice_set',
                     queryset=UserService.objects.select_related('service', 'statut')),
            Prefetch('useragence_set',
                     queryset=UserAgence.objects.select_related('agence', 'statut')),
            Prefetch('userposte_set',
                     queryset=UserPoste.objects.select_related('poste', 'statut')),
            Prefetch('usercategoriesalarie_set',
                     queryset=UserCategoriesalarie.objects.select_related('categoriesalarie', 'statut')),
            Prefetch('userallergie_set',
                     queryset=UserAllergie.objects.all()),
            'groups',
        ).order_by('id')

    def list(self, request, *args, **kwargs):
        """
        Override list() pour forcer le prefetch_related sur les objets
        de la page courante APRÈS la pagination.
        C'est la seule façon fiable de garantir le prefetch avec DRF.
        """
        from django.db.models import Prefetch

        # 1. Appliquer filtres et recherche sur le queryset de base (sans prefetch)
        base_qs = User.objects.select_related('statut').order_by('id')
        base_qs = self.filter_queryset(base_qs)

        # 2. Paginer sur le queryset de base (léger, sans prefetch)
        page_qs = self.paginate_queryset(base_qs)

        if page_qs is not None:
            # 3. Récupérer les IDs de la page courante
            user_ids = [u.id for u in page_qs]

            # 4. Recharger UNIQUEMENT ces users avec tous les prefetch
            users_avec_prefetch = list(
                User.objects.filter(id__in=user_ids)
                .select_related('statut')
                .prefetch_related(
                    Prefetch('userservice_set',
                             queryset=UserService.objects.select_related('service', 'statut')),
                    Prefetch('useragence_set',
                             queryset=UserAgence.objects.select_related('agence', 'statut')),
                    Prefetch('userposte_set',
                             queryset=UserPoste.objects.select_related('poste', 'statut')),
                    Prefetch('usercategoriesalarie_set',
                             queryset=UserCategoriesalarie.objects.select_related('categoriesalarie', 'statut')),
                    Prefetch('userallergie_set',
                             queryset=UserAllergie.objects.all()),
                    'groups',
                )
                .order_by('id')
            )

            serializer = self.get_serializer(users_avec_prefetch, many=True)
            return self.get_paginated_response(serializer.data)

        # Fallback sans pagination
        qs = self.filter_queryset(self.get_queryset())
        serializer = self.get_serializer(qs, many=True)
        return Response(serializer.data)

    @action(detail=False, methods=['post'], permission_classes=[AllowAny])
    def register(self, request):
        serializer = self.get_serializer(data=request.data)
        if serializer.is_valid():
            serializer.save()
            return Response(
                {"message": "Utilisateur créé avec succès"},
                status=status.HTTP_201_CREATED
            )
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)

    @action(detail=False, methods=['get'])
    def me(self, request):
        serializer = self.get_serializer(request.user)
        return Response(serializer.data)

    @action(detail=True, methods=['get'])
    def connection_history(self, request, pk=None):
        user = self.get_object()
        return Response({
            'username':             user.username,
            'derniere_connexion':   user.last_login.strftime("%Y-%m-%d %H:%M:%S") if user.last_login else None,
            'derniere_deconnexion': user.last_logout.strftime("%Y-%m-%d %H:%M:%S") if user.last_logout else None,
            'statut_actuel':        'Connecté' if user.is_active else 'Déconnecté'
        })

    @action(detail=True, methods=['patch'], url_path='activer')
    def activer(self, request, pk=None):
        user = self.get_object()
        user.is_active = True
        user.save()
        return Response({'message': f'Utilisateur {user.username} activé avec succès.'})

    @action(detail=True, methods=['patch'], url_path='desactiver')
    def desactiver(self, request, pk=None):
        user = self.get_object()
        user.is_active = False
        user.save()
        return Response({'message': f'Utilisateur {user.username} désactivé avec succès.'})

    @action(detail=False, methods=['patch'], url_path='activer-multiple')
    def activer_multiple(self, request):
        ids = request.data.get('ids', [])
        if not ids:
            return Response({'error': 'Aucun ID fourni.'}, status=400)
        User.objects.filter(id__in=ids).update(is_active=True)
        return Response({'message': f'{len(ids)} utilisateur(s) activé(s) avec succès.'})

    @action(detail=False, methods=['patch'], url_path='desactiver-multiple')
    def desactiver_multiple(self, request):
        ids = request.data.get('ids', [])
        if not ids:
            return Response({'error': 'Aucun ID fourni.'}, status=400)
        User.objects.filter(id__in=ids).update(is_active=False)
        return Response({'message': f'{len(ids)} utilisateur(s) désactivé(s) avec succès.'})

    @action(
        detail=False, methods=['post'],
        url_path='forgot-password',
        permission_classes=[AllowAny],
    )
    def forgot_password(self, request):
        """
        POST /api/api/users/forgot-password/
        Vérifie que le username existe, récupère le contact,
        génère un code OTP et l'envoie par SMS via AllMySMS.
        """
        username = request.data.get('username', '').strip()
        if not username:
            return Response(
                {'error': 'Le nom d\'utilisateur est requis.'},
                status=status.HTTP_400_BAD_REQUEST)

        try:
            user = User.objects.get(username=username)
        except User.DoesNotExist:
            return Response(
                {'error': 'Aucun compte trouvé avec cet identifiant.'},
                status=status.HTTP_404_NOT_FOUND)

        contact = getattr(user, 'contact', None)
        if not contact:
            return Response(
                {'error': 'Aucun numéro de contact associé à ce compte. Contactez l\'administrateur.'},
                status=status.HTTP_400_BAD_REQUEST)

        # ── Générer un code OTP à 6 chiffres ─────────────────
        code = ''.join(random.choices(string.digits, k=6))

        # ── Stocker le code en cache (10 minutes) ─────────────
        cache_key = f'reset_code_{username}'
        cache.set(cache_key, code, timeout=600)

        # ── Envoyer le SMS via AllMySMS ───────────────────────
        try:
            sms_response = requests.post(
                'https://api.allmysms.com/sms/send',
                json={
                    'login':    'gbohip',   # ← à configurer
                    'apiKey':   'd4e43a902880b7b', # ← à configurer
                    'from':     'UbiFood',
                    'to':       contact,
                    'text':     f'UbiFood - Votre code de réinitialisation : {code}\nValable 10 minutes.',
                },
                timeout=10,
            )
            if sms_response.status_code not in [200, 201]:
                return Response(
                    {'error': 'Erreur lors de l\'envoi du SMS. Réessayez.'},
                    status=status.HTTP_500_INTERNAL_SERVER_ERROR)
        except Exception:
            return Response(
                {'error': 'Service SMS indisponible. Réessayez plus tard.'},
                status=status.HTTP_503_SERVICE_UNAVAILABLE)

        # Masquer partiellement le numéro pour la réponse
        contact_masque = contact[-4:].rjust(len(contact), '*')

        return Response({
            'message':        'Code envoyé par SMS.',
            'contact_masque': contact_masque,
        }, status=status.HTTP_200_OK)


    @action(
        detail=False, methods=['post'],
        url_path='verify-code',
        permission_classes=[AllowAny],
    )
    def verify_code(self, request):
        """
        POST /api/api/users/verify-code/
        Vérifie le code OTP saisi par l'utilisateur.
        """
        username = request.data.get('username', '').strip()
        code     = request.data.get('code', '').strip()

        if not username or not code:
            return Response(
                {'error': 'Identifiant et code requis.'},
                status=status.HTTP_400_BAD_REQUEST)

        cache_key   = f'reset_code_{username}'
        code_stocke = cache.get(cache_key)

        if not code_stocke:
            return Response(
                {'error': 'Code expiré ou invalide. Recommencez.'},
                status=status.HTTP_400_BAD_REQUEST)

        if code_stocke != code:
            return Response(
                {'error': 'Code incorrect.'},
                status=status.HTTP_400_BAD_REQUEST)

        # Générer un token temporaire pour la réinitialisation
        reset_token = ''.join(random.choices(string.ascii_letters + string.digits, k=32))
        cache.set(f'reset_token_{username}', reset_token, timeout=600)
        cache.delete(cache_key)  # Supprimer le code utilisé

        return Response({
            'message':     'Code vérifié.',
            'reset_token': reset_token,
        }, status=status.HTTP_200_OK)


    @action(
        detail=False, methods=['post'],
        url_path='reset-password',
        permission_classes=[AllowAny],
    )
    
    def reset_password(self, request):
        """
        POST /api/api/users/reset-password/
        Réinitialise le mot de passe avec le token temporaire.
        Pas besoin de l'ancien mot de passe.
        """
        username    = request.data.get('username', '').strip()
        reset_token = request.data.get('reset_token', '').strip()
        nouveau_mdp = request.data.get('password', '').strip()

        if not username or not reset_token or not nouveau_mdp:
            return Response(
                {'error': 'Données incomplètes.'},
                status=status.HTTP_400_BAD_REQUEST)

        # Vérifier le token
        token_stocke = cache.get(f'reset_token_{username}')
        if not token_stocke or token_stocke != reset_token:
            return Response(
                {'error': 'Session expirée. Recommencez depuis le début.'},
                status=status.HTTP_400_BAD_REQUEST)

        try:
            user = User.objects.get(username=username)
        except User.DoesNotExist:
            return Response(
                {'error': 'Utilisateur introuvable.'},
                status=status.HTTP_404_NOT_FOUND)

        # Valider la complexité
        try:
            validate_password_custom(nouveau_mdp)
        except DjangoValidationError as e:
            return Response(
                {'error': e.messages},
                status=status.HTTP_400_BAD_REQUEST)

        user.set_password(nouveau_mdp)
        user.save()
        cache.delete(f'reset_token_{username}')

        return Response(
            {'message': 'Mot de passe réinitialisé avec succès.'},
            status=status.HTTP_200_OK)
    
    @action(detail=True, methods=['patch'], url_path='changer-mot-de-passe')
    def changer_mot_de_passe(self, request, pk=None):
        """
        PATCH /api/api/users/{id}/changer-mot-de-passe/
        - User normal : peut uniquement modifier SON propre mot de passe
        - super_admin / admin : peuvent modifier celui de n'importe qui
        Règles : 8 caractères min | 1 majuscule | 1 chiffre | 1 caractère spécial
        """
        user        = self.get_object()
        nouveau_mdp = request.data.get('password')
        old_password = request.data.get('old_password')

        # ✅ Vérifier que l'user ne modifie que son propre compte
        # sauf si c'est un admin/super_admin
        groupes_user = request.user.groups.values_list('name', flat=True)
        is_admin = (
            request.user.is_superuser or
            'super_admin' in groupes_user or
            'admin' in groupes_user
        )

        if not is_admin and request.user.id != user.id:
            return Response(
                {'error': 'Vous ne pouvez modifier que votre propre mot de passe.'},
                status=status.HTTP_403_FORBIDDEN
            )

        # ✅ Si ce n'est pas un admin, vérifier l'ancien mot de passe
        if not is_admin:
            if not old_password:
                return Response(
                    {'error': 'L\'ancien mot de passe est requis.'},
                    status=status.HTTP_400_BAD_REQUEST
                )
            if not user.check_password(old_password):
                return Response(
                    {'error': 'Mot de passe actuel incorrect.'},
                    status=status.HTTP_400_BAD_REQUEST
                )

        if not nouveau_mdp:
            return Response(
                {'error': 'Le nouveau mot de passe est requis.'},
                status=status.HTTP_400_BAD_REQUEST
            )

        # ✅ Validation de la complexité
        try:
            validate_password_custom(nouveau_mdp)
        except DjangoValidationError as e:
            return Response(
                {'error': e.messages},
                status=status.HTTP_400_BAD_REQUEST
            )

        # ✅ Hachage et sauvegarde
        user.set_password(nouveau_mdp)
        user.save()
        return Response({'message': 'Mot de passe modifié avec succès.'})

class AllUserNopginListAPIView(ListAPIView):
    queryset         = User.objects.all()
    serializer_class = UserSerializer
    permission_classes = [IsAuthenticated]
    pagination_class = None

    
class GroupViewSet(viewsets.ModelViewSet):
    """
    GroupViewSet n'hérite pas de BaseViewSet car Group est un modèle Django natif
    qui n'a pas les champs user_created / user_updated de BaseModel.
    """
    queryset = Group.objects.all()
    serializer_class = GroupSerializer

    @action(detail=False, methods=['delete'])
    def bulk_delete(self, request):
        ids = request.data.get('ids', [])

        if not isinstance(ids, list) or not ids:
            return Response(
                {"error": "Veuillez fournir une liste d'IDs valide et non vide"},
                status=status.HTTP_400_BAD_REQUEST
            )

        if not all(isinstance(i, int) for i in ids):
            return Response(
                {"error": "Tous les IDs doivent être des entiers"},
                status=status.HTTP_400_BAD_REQUEST
            )

        deleted_count, _ = Group.objects.filter(id__in=ids).delete()

        return Response({
            "message": f"{deleted_count} rôle(s) supprimé(s) avec succès",
            "deleted_count": deleted_count
        })


# ============================================================
# RÉFÉRENTIELS
# ============================================================

class AgenceViewSet(BaseViewSet):
    queryset = Agence.objects.all().order_by('id')
    serializer_class = AgenceSerializer
    filterset_fields = ['nom_agence']
    search_fields = ['nom_agence']
    ordering_fields = ['nom_agence', 'created_at']


class AllAgenceNopginListAPIView(ListAPIView):
    queryset = Agence.objects.all()
    serializer_class = AgenceSerializer
    permission_classes = [AllowAny]
    pagination_class = None


class EtatViewSet(BaseViewSet):
    queryset = Etat.objects.all().order_by('id')
    serializer_class = EtatSerializer


class StatutViewSet(BaseViewSet):
    queryset = Statut.objects.all().order_by('id')
    serializer_class = StatutSerializer


class PublicStatutListAPIView(ListAPIView):
    queryset = Statut.objects.all()
    serializer_class = StatutSerializer
    permission_classes = [AllowAny]
    pagination_class = None


class PublicEtatListAPIView(ListAPIView):
    queryset = Etat.objects.all()
    serializer_class = EtatSerializer
    permission_classes = [AllowAny]
    pagination_class = None


class ServiceViewSet(BaseViewSet):
    queryset = Service.objects.all().order_by('id')
    serializer_class = ServiceSerializer


class TypeBesoinViewSet(BaseViewSet):
    queryset = TypeBesoin.objects.all().order_by('id')
    serializer_class = TypeBesoinSerializer


class PrioriteViewSet(BaseViewSet):
    queryset = Priorite.objects.all().order_by('id')
    serializer_class = PrioriteSerializer


class UserServiceViewSet(BaseViewSet):
    queryset = (
        UserService.objects
        .select_related('user', 'service', 'statut')
        .order_by('id')
    )
    serializer_class = UserServiceSerializer
    filterset_fields = ['user', 'service', 'statut']  # ✅ AJOUTÉ

    def perform_update(self, serializer):
        instance = serializer.instance
        new_user = serializer.validated_data.get('user', instance.user)
        if instance.user_id != new_user.id:
            from rest_framework.exceptions import PermissionDenied
            raise PermissionDenied(
                "Impossible de changer le user d'un UserService existant."
            )
        super().perform_update(serializer)


class UserAgenceViewSet(BaseViewSet):
    queryset = (
        UserAgence.objects
        .select_related('user', 'agence', 'statut')
        .order_by('id')
    )
    serializer_class = UserAgenceSerializer
    filterset_fields = ['user', 'agence', 'statut']  # ✅ AJOUTÉ

    def perform_update(self, serializer):
        instance = serializer.instance
        new_user = serializer.validated_data.get('user', instance.user)
        if instance.user_id != new_user.id:
            from rest_framework.exceptions import PermissionDenied
            raise PermissionDenied(
                "Impossible de changer le user d'un UserAgence existant."
            )
        super().perform_update(serializer)


class FonctionViewSet(BaseViewSet):
    queryset = Fonction.objects.all().order_by('id')
    serializer_class = FonctionSerializer
    filterset_fields = ['libelle']
    search_fields = ['libelle']
    ordering_fields = ['libelle', 'created_at']


class CategoriesalarieViewSet(BaseViewSet):
    queryset = Categoriesalarie.objects.all().order_by('id')
    serializer_class = CategoriesalarieSerializer
    filterset_fields = ['libelle']
    search_fields = ['libelle']
    ordering_fields = ['libelle', 'created_at']


class PrestataireViewSet(BaseViewSet):
    queryset = Prestataire.objects.all().order_by('id')
    serializer_class = PrestataireSerializer
    filterset_fields = ['libelle']
    search_fields = ['libelle']
    ordering_fields = ['libelle', 'created_at']


class PlatPrestataireViewSet(BaseViewSet):
    queryset = PlatPrestataire.objects.all().order_by('-date_debut')
    serializer_class = PlatPrestataireSerializer
    filterset_fields = ['prestataire', 'statut']
    ordering_fields  = ['date_debut', 'created_at']


class AgencePrestataireViewSet(BaseViewSet):
    queryset = AgencePrestataire.objects.all().order_by('-date_debut')
    serializer_class = AgencePrestataireSerializer
    filterset_fields = ['prestataire', 'agence', 'statut']
    ordering_fields  = ['date_debut', 'created_at']


class PlatCategoriesalarieViewSet(BaseViewSet):
    queryset = PlatCategoriesalarie.objects.all().order_by('-date_debut')
    serializer_class = PlatCategoriesalarieSerializer
    filterset_fields = ['categoriesalarie', 'statut']
    ordering_fields  = ['date_debut', 'created_at']


class PosteViewSet(BaseViewSet):
    queryset = Poste.objects.select_related('fonction', 'service').all().order_by('id')
    serializer_class = PosteSerializer
    filterset_fields = ['libelle', 'fonction', 'service']
    search_fields    = ['libelle']
    ordering_fields  = ['libelle', 'created_at']


# ============================================================
# BESOINS
# ============================================================

class BesoinViewSet(BaseViewSet):
    queryset = (
        Besoin.objects
        .select_related('typebesoin', 'etat', 'priorite', 'user')
        .prefetch_related('besoinservice_set__service', 'documents')
        .order_by('id')
    )
    serializer_class  = BesoinSerializer
    filterset_fields  = ['reference', 'titre', 'typebesoin', 'etat', 'priorite', 'user']
    search_fields     = ['reference', 'titre', 'description']
    ordering_fields   = ['date_debut', 'date_fin', 'created_at']

    # ── Mes besoins ───────────────────────────────────────────
    @action(detail=False, methods=['get'])
    def mes_besoins(self, request):
        besoins    = self.queryset.filter(user=request.user)
        page       = self.paginate_queryset(besoins)
        serializer = self.get_serializer(page, many=True)
        return self.get_paginated_response(serializer.data)

    # ── Statistiques par état ─────────────────────────────────
    @action(detail=False, methods=['get'], url_path='statistiques-par-etat')
    def statistiques_par_etat(self, request):
        data = (
            self.queryset
            .values('etat__libelle_etat')
            .annotate(nombre=Count('id'))
            .order_by('etat__libelle_etat')
        )
        result = [{'etat': d['etat__libelle_etat'], 'nombre': d['nombre']} for d in data]
        return Response(result)

    # ── Statistiques par priorité ─────────────────────────────
    @action(detail=False, methods=['get'], url_path='statistiques-par-priorite')
    def statistiques_par_priorite(self, request):
        data = (
            self.queryset
            .values('priorite__libelle')
            .annotate(nombre=Count('id'))
            .order_by('priorite__libelle')
        )
        result = [{'priorite': d['priorite__libelle'], 'nombre': d['nombre']} for d in data]
        return Response(result)

    # ── Statistiques globales ─────────────────────────────────
    @action(detail=False, methods=['get'], url_path='statistiques-globales')
    def statistiques_globales(self, request):
        etats = (
            self.queryset
            .values('etat__libelle_etat')
            .annotate(nombre=Count('id'))
            .order_by('etat__libelle_etat')
        )
        etat_result = [{'etat': d['etat__libelle_etat'], 'nombre': d['nombre']} for d in etats]

        priorites = (
            self.queryset
            .values('priorite__libelle')
            .annotate(nombre=Count('id'))
            .order_by('priorite__libelle')
        )
        priorite_result = [{'priorite': d['priorite__libelle'], 'nombre': d['nombre']} for d in priorites]

        total_besoins     = sum(item['nombre'] for item in etat_result)
        total_cloture     = sum(item['nombre'] for item in etat_result if item['etat'].lower() in ['clôturé', 'cloturé', 'clôturés'])
        total_annule      = sum(item['nombre'] for item in etat_result if item['etat'].lower() in ['annulé', 'annulés'])
        total_non_cloture = total_besoins - total_cloture - total_annule

        return Response({
            'par_etat':          etat_result,
            'par_priorite':      priorite_result,
            'total_non_cloture': total_non_cloture,
            'total_besoin':      total_besoins,
        })

    # ── Mise à jour ───────────────────────────────────────────
    def update(self, request, *args, **kwargs):
        partial    = kwargs.pop('partial', False)
        instance   = self.get_object()
        serializer = self.get_serializer(instance, data=request.data, partial=partial)

        if not serializer.is_valid():
            return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)

        self.perform_update(serializer)
        return Response(serializer.data)

    # ── ✅ Ajouter des documents à un besoin existant ─────────
    # POST /api/api/besoins/{id}/ajouter-documents/
    # Champ attendu : documents[] (un ou plusieurs fichiers)
    @action(
        detail=True,
        methods=['post'],
        url_path='ajouter-documents',
        parser_classes=[MultiPartParser, FormParser],
    )
    def ajouter_documents(self, request, pk=None):
        besoin = self.get_object()
        files  = request.FILES.getlist('documents[]')

        if not files:
            return Response(
                {'error': 'Aucun fichier fourni. Utilisez la clé documents[].'},
                status=status.HTTP_400_BAD_REQUEST
            )

        created = []
        for file in files:
            doc = BesoinDocument.objects.create(besoin=besoin, document=file)
            created.append(BesoinDocumentSerializer(doc).data)

        return Response(created, status=status.HTTP_201_CREATED)

    # ── ✅ Supprimer un document rattaché à un besoin ─────────
    # DELETE /api/api/besoins/{id}/supprimer-document/?document_id=<id>
    @action(
        detail=True,
        methods=['delete'],
        url_path='supprimer-document',
    )
    def supprimer_document(self, request, pk=None):
        besoin      = self.get_object()
        document_id = request.query_params.get('document_id')

        if not document_id:
            return Response(
                {'error': 'Paramètre document_id requis.'},
                status=status.HTTP_400_BAD_REQUEST
            )

        try:
            doc = BesoinDocument.objects.get(id=document_id, besoin=besoin)
        except BesoinDocument.DoesNotExist:
            return Response(
                {'error': 'Document introuvable pour ce besoin.'},
                status=status.HTTP_404_NOT_FOUND
            )

        # ✅ Supprimer aussi le fichier physique du disque
        if doc.document:
            try:
                doc.document.delete(save=False)
            except Exception:
                pass  # Ne pas bloquer si le fichier est déjà absent

        doc.delete()
        return Response(
            {'message': 'Document supprimé avec succès.'},
            status=status.HTTP_200_OK
        )


class BesoinServiceViewSet(BaseViewSet):
    queryset = BesoinService.objects.all().order_by('id')
    serializer_class = BesoinServiceSerializer


class BesoinDocumentViewSet(BaseViewSet):
    queryset = BesoinDocument.objects.all().order_by('id')
    serializer_class = BesoinDocumentSerializer


class PublicTypeBesoinListAPIView(ListAPIView):
    queryset = TypeBesoin.objects.all()
    serializer_class = TypeBesoinSerializer
    permission_classes = [AllowAny]
    pagination_class = None


class PublicPrioriteListAPIView(ListAPIView):
    queryset = Priorite.objects.all()
    serializer_class = PrioriteSerializer
    permission_classes = [AllowAny]
    pagination_class = None


class BesoinCreateView(APIView):
    permission_classes = [AllowAny]
    parser_classes = [MultiPartParser, FormParser]

    def post(self, request, format=None):
        serializer = BesoinWithDocumentsSerializer(data=request.data, context={'request': request})
        if serializer.is_valid():
            besoin = serializer.save()
            return Response(
                {"message": "Besoin et documents enregistrés avec succès", "besoin_id": besoin.id},
                status=status.HTTP_201_CREATED
            )
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)


# ============================================================
# CANTINE
# ============================================================

class TypePlatViewSet(BaseViewSet):
    queryset = TypePlat.objects.all().order_by('id')
    serializer_class = TypePlatSerializer


class TypeEquipeViewSet(BaseViewSet):
    queryset = TypeEquipe.objects.all().order_by('id')
    serializer_class = TypeEquipeSerializer


class PublicTypePlatListAPIView(ListAPIView):
    queryset = TypePlat.objects.all()
    serializer_class = TypePlatSerializer
    permission_classes = [AllowAny]
    pagination_class = None


class PublicTypeEquipeListAPIView(ListAPIView):
    queryset = TypeEquipe.objects.all()
    serializer_class = TypeEquipeSerializer
    permission_classes = [AllowAny]
    pagination_class = None


class PlatViewSet(BaseViewSet):
    queryset = (
        Plat.objects
        .select_related('type_plat', 'agence')
        .prefetch_related('images')
        .order_by('id')
    )
    serializer_class = PlatSerializer
    filterset_fields = ['type_plat']
    search_fields = ['nom', 'description']
    parser_classes = [MultiPartParser, FormParser, JSONParser]

    def create(self, request, *args, **kwargs):
        serializer = self.get_serializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        plat = serializer.save(
            user_created=request.user,
            user_updated=request.user
        )
        new_images = request.FILES.getlist('new_images')
        for i, image in enumerate(new_images):
            PlatImage.objects.create(plat=plat, image=image, is_principale=(i == 0))
        output_serializer = self.get_serializer(plat)
        return Response(output_serializer.data, status=status.HTTP_201_CREATED)

    def partial_update(self, request, *args, **kwargs):
        plat = self.get_object()
        data = request.data.copy()

        images_to_delete = request.data.getlist('images_to_delete[]', [])
        for img_id in images_to_delete:
            try:
                PlatImage.objects.get(id=img_id, plat=plat).delete()
            except PlatImage.DoesNotExist:
                pass

        new_images = request.FILES.getlist('new_images')
        for image in new_images:
            PlatImage.objects.create(plat=plat, image=image)

        serializer = self.get_serializer(plat, data=data, partial=True)
        serializer.is_valid(raise_exception=True)
        self.perform_update(serializer)

        output_serializer = self.get_serializer(plat)
        return Response(output_serializer.data, status=status.HTTP_200_OK)


class PlatImageViewSet(BaseViewSet):
    queryset = PlatImage.objects.all().order_by('id')
    serializer_class = PlatImageSerializer


class PlatCreateView(APIView):
    parser_classes = [MultiPartParser, FormParser]
    permission_classes = [IsAuthenticated]

    def post(self, request, *args, **kwargs):
        serializer = PlatCreateWithImagesSerializer(data=request.data, context={'request': request})
        if serializer.is_valid():
            plat = serializer.save(user_created=request.user, user_updated=request.user)
            return Response({'message': 'Plat enregistré', 'id': plat.id}, status=status.HTTP_201_CREATED)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)


class PlatUploadAPIView(APIView):
    parser_classes = [MultiPartParser, FormParser]
    permission_classes = [IsAuthenticated]

    def post(self, request, *args, **kwargs):
        fichiers = request.FILES.getlist('images[]')
        data = request.data.copy()
        data.setlist('images', [])
        for fichier in fichiers:
            data.appendlist('images', {'image': fichier, 'is_principale': False})

        serializer = PlatUploadSerializer(data=data, context={'request': request})
        if serializer.is_valid():
            plat = serializer.save()
            return Response({'message': 'Plat créé avec succès', 'plat_id': plat.id}, status=status.HTTP_201_CREATED)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)


class MenuViewSet(BaseViewSet):
    queryset = (
        Menu.objects
        .select_related('agence', 'typeequipe')
        .prefetch_related(
            'menu_plats__plat__images',
            'menu_plats__plat__type_plat',
            'menu_plats__plat__agence',
        )
        .order_by('-date_menu')
    )
    serializer_class = MenuSerializer
    filterset_fields = {
        'date_menu':  ['exact', 'gte', 'lte'],
        'agence':     ['exact'],
        'typeequipe': ['exact'],
    }


class MenuPlatViewSet(BaseViewSet):
    queryset = MenuPlat.objects.select_related('plat').all().order_by('id')
    serializer_class = MenuPlatSerializer
    filterset_fields = ['menu']


# ─────────────────────────────────────────────────────────────────────────────
# COMMANDE
# ─────────────────────────────────────────────────────────────────────────────

class CommandeViewSet(BaseViewSet):
    serializer_class = CommandeSerializer
    filterset_fields = {
        'menu':            ['exact'],
        'user':            ['exact'],
        'statut':          ['exact'],
        'menu__agence':    ['exact'],
        'menu__date_menu': ['exact', 'gte', 'lte'],
    }

    def get_queryset(self):
        return (
            Commande.objects
            .select_related('user', 'menu__agence', 'menu__typeequipe', 'plat')
            .prefetch_related('plat__images')
            .order_by('-date_commande')
        )

    def create(self, request, *args, **kwargs):
        serializer = self.get_serializer(data=request.data)
        serializer.is_valid(raise_exception=True)

        menu = serializer.validated_data['menu']
        plat = serializer.validated_data['plat']
        user = request.user

        if timezone.now() > _get_deadline(menu.date_menu):
            from rest_framework.exceptions import ValidationError
            raise ValidationError(
                "Impossible de commander : le délai de 48h avant le menu est dépassé."
            )

        commande_jour = Commande.objects.filter(
            user=user,
            menu__date_menu=menu.date_menu,
            statut__in=['en_attente', 'retiree'],
        ).exclude(menu=menu, plat=plat)

        if commande_jour.exists():
            from rest_framework.exceptions import ValidationError
            raise ValidationError(
                f"Vous avez déjà une commande pour le {menu.date_menu}. "
                "Une seule commande est autorisée par jour."
            )

        commande_existante = Commande.objects.filter(
            user=user, menu=menu, plat=plat, statut='annulee'
        ).first()

        if commande_existante:
            commande_existante.statut          = 'en_attente'
            commande_existante.date_annulation = None
            commande_existante.user_updated    = user
            commande_existante.save()
            output = CommandeSerializer(commande_existante, context={'request': request})
            return Response(output.data, status=status.HTTP_200_OK)

        serializer.save(user=user, user_created=user, user_updated=user)
        return Response(serializer.data, status=status.HTTP_201_CREATED)

    @action(detail=True, methods=['patch'], url_path='annuler')
    def annuler(self, request, pk=None):
        commande   = self.get_object()
        serializer = CommandeAnnulationSerializer(
            commande, data={}, partial=True, context={'request': request}
        )
        serializer.is_valid(raise_exception=True)
        serializer.save()
        return Response(
            CommandeSerializer(commande, context={'request': request}).data,
            status=status.HTTP_200_OK
        )

    @action(detail=False, methods=['get'], url_path='mes-commandes')
    def mes_commandes(self, request):
        qs = self.get_queryset().filter(user=request.user)
        qs = self.filter_queryset(qs)
        serializer = self.get_serializer(qs, many=True)
        return Response(serializer.data)

    @action(detail=False, methods=['get'], url_path='par-menu')
    def par_menu(self, request):
        menu_id = request.query_params.get('menu')
        if not menu_id:
            return Response({'detail': 'Paramètre menu requis.'}, status=400)

        commandes = self.get_queryset().filter(menu_id=menu_id)
        grouped   = {}
        for cmd in commandes:
            uid = cmd.user_id
            if uid not in grouped:
                grouped[uid] = {
                    'user_id':   uid,
                    'user_nom':  f"{cmd.user.first_name} {cmd.user.last_name}".strip() or cmd.user.username,
                    'commandes': [],
                    'a_retire':  Retrait.objects.filter(user_id=uid, menu_id=menu_id).exists(),
                }
            grouped[uid]['commandes'].append(
                CommandeSerializer(cmd, context={'request': request}).data
            )
        return Response(list(grouped.values()))

    @action(detail=False, methods=['get'], url_path='par-agence-periode')
    def par_agence_periode(self, request):
        date_debut = request.query_params.get('date_debut')
        date_fin   = request.query_params.get('date_fin')
        agence     = request.query_params.get('agence')
        typeequipe = request.query_params.get('typeequipe')
        statut     = request.query_params.get('statut')

        qs = self.get_queryset()
        if date_debut:  qs = qs.filter(menu__date_menu__gte=date_debut)
        if date_fin:    qs = qs.filter(menu__date_menu__lte=date_fin)
        if agence:      qs = qs.filter(menu__agence__id=agence)
        if typeequipe:  qs = qs.filter(menu__typeequipe__id=typeequipe)
        if statut:      qs = qs.filter(statut=statut)

        serializer = CommandeSerializer(qs, many=True, context={'request': request})
        return Response(serializer.data)

    @action(detail=False, methods=['get'], url_path='recherche-par-badge')
    def recherche_par_badge(self, request):
        badge   = request.query_params.get('badge', '').strip()
        menu_id = request.query_params.get('menu')

        if not badge:   return Response({'detail': 'Badge requis.'}, status=400)
        if not menu_id: return Response({'detail': 'Menu requis.'}, status=400)

        try:
            user = User.objects.get(username=badge)
        except User.DoesNotExist:
            return Response({'detail': f'Aucun agent trouvé avec le badge "{badge}".'}, status=404)

        a_retire  = Retrait.objects.filter(user=user, menu_id=menu_id).exists()
        commandes = self.get_queryset().filter(user=user, menu_id=menu_id)

        return Response({
            'user_id':   user.id,
            'user_nom':  f"{user.first_name} {user.last_name}".strip() or user.username,
            'badge':     user.username,
            'a_retire':  a_retire,
            'commandes': CommandeSerializer(commandes, many=True, context={'request': request}).data,
        })

class DashboardCantineView(APIView):
    """
    GET /api/api/dashboard-cantine/
    Params:
      date_debut   : YYYY-MM-DD
      date_fin     : YYYY-MM-DD
      agences      : "1,2,3"
      typeequipes  : "1,2"
      statuts      : "en_attente,retiree,annulee"
    """
    permission_classes = [IsAuthenticated]

    def get(self, request):
        from django.db.models import Count, Q, Sum, F, DecimalField
        from django.db.models import ExpressionWrapper
        from django.db.models.functions import ExtractMonth, Coalesce
        from decimal import Decimal
        from datetime import date as date_type, timedelta

        # ── Paramètres ────────────────────────────────────────
        date_debut  = request.query_params.get('date_debut',  None)
        date_fin    = request.query_params.get('date_fin',    None)
        agences_str = request.query_params.get('agences',     None)
        equipes_str = request.query_params.get('typeequipes', None)
        statuts_str = request.query_params.get('statuts',     None)

        agences_ids = [
            int(x) for x in agences_str.split(',') if x.strip().isdigit()
        ] if agences_str else []

        equipes_ids = [
            int(x) for x in equipes_str.split(',') if x.strip().isdigit()
        ] if equipes_str else []

        statuts_list = [
            x.strip() for x in statuts_str.split(',')
            if x.strip() in ['en_attente', 'retiree', 'annulee']
        ] if statuts_str else []

        # ── Querysets de base ─────────────────────────────────
        commande_qs = Commande.objects.all()
        menu_qs     = Menu.objects.all()

        if date_debut:
            commande_qs = commande_qs.filter(menu__date_menu__gte=date_debut)
            menu_qs     = menu_qs.filter(date_menu__gte=date_debut)
        if date_fin:
            commande_qs = commande_qs.filter(menu__date_menu__lte=date_fin)
            menu_qs     = menu_qs.filter(date_menu__lte=date_fin)
        if agences_ids:
            commande_qs = commande_qs.filter(menu__agence__id__in=agences_ids)
            menu_qs     = menu_qs.filter(agence__id__in=agences_ids)
        if equipes_ids:
            commande_qs = commande_qs.filter(menu__typeequipe__id__in=equipes_ids)
            menu_qs     = menu_qs.filter(typeequipe__id__in=equipes_ids)
        if statuts_list:
            commande_qs = commande_qs.filter(statut__in=statuts_list)

        # ── KPIs globaux ──────────────────────────────────────
        total_commandes  = commande_qs.count()
        total_en_attente = commande_qs.filter(statut='en_attente').count()
        total_retirees   = commande_qs.filter(statut='retiree').count()
        total_annulees   = commande_qs.filter(statut='annulee').count()
        total_menus      = menu_qs.count()
        taux_retrait     = round((total_retirees / total_commandes * 100) if total_commandes > 0 else 0, 1)
        taux_annulation  = round((total_annulees / total_commandes * 100) if total_commandes > 0 else 0, 1)

        # ── Par mois ──────────────────────────────────────────
        commandes_par_mois = (
            commande_qs
            .annotate(mois=ExtractMonth('menu__date_menu'))
            .values('mois')
            .annotate(
                total=Count('id'),
                en_attente=Count('id', filter=Q(statut='en_attente')),
                retiree=Count('id',    filter=Q(statut='retiree')),
                annulee=Count('id',    filter=Q(statut='annulee')),
            )
            .order_by('mois')
        )
        par_mois = {i: {'total': 0, 'en_attente': 0, 'retiree': 0, 'annulee': 0} for i in range(1, 13)}
        for row in commandes_par_mois:
            par_mois[row['mois']] = {
                'total':      int(row['total']      or 0),
                'en_attente': int(row['en_attente'] or 0),
                'retiree':    int(row['retiree']    or 0),
                'annulee':    int(row['annulee']    or 0),
            }

        # ── Top plats ─────────────────────────────────────────
        top_plats = (
            commande_qs.values('plat__nom', 'plat__type_plat__libelle')
            .annotate(nombre=Count('id')).order_by('-nombre')[:10]
        )
        top_plats_data = [
            {
                'nom':       str(row['plat__nom'] or ''),
                'type_plat': str(row['plat__type_plat__libelle'] or ''),
                'nombre':    int(row['nombre'] or 0),
            }
            for row in top_plats
        ]

        # ── Par agence ────────────────────────────────────────
        par_agence = (
            commande_qs.values('menu__agence__nom_agence')
            .annotate(
                total=Count('id'),
                en_attente=Count('id', filter=Q(statut='en_attente')),
                retiree=Count('id',    filter=Q(statut='retiree')),
                annulee=Count('id',    filter=Q(statut='annulee')),
            ).order_by('-total')
        )
        par_agence_data = [
            {
                'agence':     str(row['menu__agence__nom_agence'] or ''),
                'total':      int(row['total']      or 0),
                'en_attente': int(row['en_attente'] or 0),
                'retiree':    int(row['retiree']    or 0),
                'annulee':    int(row['annulee']    or 0),
            }
            for row in par_agence
        ]

        # ── Par type équipe ───────────────────────────────────
        par_typeequipe_data = [
            {
                'typeequipe': str(row['menu__typeequipe__libelle'] or ''),
                'total':      int(row['total'] or 0),
            }
            for row in (
                commande_qs.values('menu__typeequipe__libelle')
                .annotate(total=Count('id')).order_by('-total')
            )
        ]

        # ── Par type plat ─────────────────────────────────────
        par_typeplat_data = [
            {
                'type_plat': str(row['plat__type_plat__libelle'] or ''),
                'total':     int(row['total'] or 0),
            }
            for row in (
                commande_qs.values('plat__type_plat__libelle')
                .annotate(total=Count('id')).order_by('-total')
            )
        ]

        # ── Menus récents ─────────────────────────────────────
        menus_recents_data = [
            {
                'id':           int(m.id),
                'date_menu':    str(m.date_menu),
                'agence':       str(m.agence.nom_agence  if m.agence     else ''),
                'typeequipe':   str(m.typeequipe.libelle if m.typeequipe else ''),
                'nb_commandes': int(m.nb_commandes or 0),
            }
            for m in (
                menu_qs.select_related('agence', 'typeequipe')
                .annotate(nb_commandes=Count('commandes'))
                .order_by('-date_menu')[:5]
            )
        ]

        # ── Évolution hebdomadaire ────────────────────────────
        today    = date_type.today()
        semaines = []
        for i in range(3, -1, -1):
            debut_s = today - timedelta(weeks=i + 1)
            fin_s   = today - timedelta(weeks=i)
            qs_s    = Commande.objects.filter(menu__date_menu__gte=debut_s, menu__date_menu__lt=fin_s)
            if agences_ids: qs_s = qs_s.filter(menu__agence__id__in=agences_ids)
            if equipes_ids: qs_s = qs_s.filter(menu__typeequipe__id__in=equipes_ids)
            if statuts_list: qs_s = qs_s.filter(statut__in=statuts_list)
            semaines.append({'label': 'S-' + str(i + 1) if i > 0 else 'Cette sem.', 'total': qs_s.count()})


        # ── Evolution des commandes par statut par date menu ────
        # Retourne chaque date_menu avec les 3 compteurs de statut
        # Filtres de dates/agences/équipes appliqués, pas de filtre statut
        # (on veut toujours les 3 statuts côte à côte)
        evo_qs = Commande.objects.all()
        if date_debut:  evo_qs = evo_qs.filter(menu__date_menu__gte=date_debut)
        if date_fin:    evo_qs = evo_qs.filter(menu__date_menu__lte=date_fin)
        if agences_ids: evo_qs = evo_qs.filter(menu__agence__id__in=agences_ids)
        if equipes_ids: evo_qs = evo_qs.filter(menu__typeequipe__id__in=equipes_ids)

        evo_par_date_qs = (
            evo_qs
            .values('menu__date_menu')
            .annotate(
                en_attente=Count('id', filter=Q(statut='en_attente')),
                retiree=Count('id',    filter=Q(statut='retiree')),
                annulee=Count('id',    filter=Q(statut='annulee')),
            )
            .order_by('menu__date_menu')
        )
        evolution_par_date = [
            {
                'date':       str(row['menu__date_menu']),
                'en_attente': int(row['en_attente'] or 0),
                'retiree':    int(row['retiree']    or 0),
                'annulee':    int(row['annulee']    or 0),
            }
            for row in evo_par_date_qs
        ]
        # ════════════════════════════════════════════════════
        # FACTURATION PRESTATAIRE
        # PlatPrestataire = tarif global par prestataire sur une période
        # Champs : id, montant, date_debut, date_fin, prestataire,
        #          statut, created_at, updated_at
        # Logique : pour chaque commande retirée, on cherche le
        # PlatPrestataire actif à la date du menu.
        # La comparaison date_menu (DateField) vs date_debut (DateTimeField)
        # nécessite de travailler avec .date() pour éviter le warning timezone.
        # ════════════════════════════════════════════════════
        commandes_retirees = commande_qs.filter(statut='retiree').select_related(
            'plat', 'menu', 'menu__agence', 'menu__typeequipe'
        )

        total_montant_prestataire = Decimal('0.00')
        fact_prest_par_mois       = {i: Decimal('0.00') for i in range(1, 13)}
        fact_prest_par_agence     = {}
        fact_prest_par_prest      = {}

        for cmd in commandes_retirees:
            date_menu = cmd.menu.date_menu  # DateField → objet date Python

            # Chercher le PlatPrestataire actif à la date du menu
            # date_debut et date_fin sont des DateTimeField → on utilise __date
            pp = PlatPrestataire.objects.filter(
                date_debut__date__lte=date_menu,
            ).filter(
                Q(date_fin__isnull=True) | Q(date_fin__date__gte=date_menu)
            ).order_by('-date_debut').first()

            if not pp:
                continue

            montant = pp.montant or Decimal('0.00')
            total_montant_prestataire += montant

            # Par mois
            mois = date_menu.month
            fact_prest_par_mois[mois] = fact_prest_par_mois.get(mois, Decimal('0.00')) + montant

            # Par agence
            agence_nom = cmd.menu.agence.nom_agence
            if agence_nom not in fact_prest_par_agence:
                fact_prest_par_agence[agence_nom] = {'agence': agence_nom, 'montant': Decimal('0.00'), 'nb_commandes': 0}
            fact_prest_par_agence[agence_nom]['montant']      += montant
            fact_prest_par_agence[agence_nom]['nb_commandes'] += 1

            # Par prestataire — accès direct aux champs du modele Prestataire
            try:
                p_obj = pp.prestataire
                if not p_obj:
                    prest_nom = 'Inconnu'
                else:
                    # Essayer les champs nom courants — adapter selon votre modele
                    prest_nom = (
                        getattr(p_obj, 'nom', None) or
                        getattr(p_obj, 'libelle', None) or
                        getattr(p_obj, 'name', None) or
                        getattr(p_obj, 'raison_sociale', None) or
                        'Prestataire ' + str(p_obj.pk)
                    )
            except Exception:
                prest_nom = 'Inconnu'
            if prest_nom not in fact_prest_par_prest:
                fact_prest_par_prest[prest_nom] = {'prestataire': prest_nom, 'montant': Decimal('0.00'), 'nb_commandes': 0}
            fact_prest_par_prest[prest_nom]['montant']      += montant
            fact_prest_par_prest[prest_nom]['nb_commandes'] += 1

        # Sérialisation Decimal → float
        fact_fournisseur = {
            'total_montant': float(total_montant_prestataire),
            'par_mois': {k: float(v) for k, v in fact_prest_par_mois.items()},
            'par_agence': [
                {'agence': v['agence'], 'montant': float(v['montant']), 'nb_commandes': v['nb_commandes']}
                for v in sorted(fact_prest_par_agence.values(), key=lambda x: x['montant'], reverse=True)
            ],
            'par_prestataire': [
                {'prestataire': v['prestataire'], 'montant': float(v['montant']), 'nb_commandes': v['nb_commandes']}
                for v in sorted(fact_prest_par_prest.values(), key=lambda x: x['montant'], reverse=True)
            ],
        }

        # ════════════════════════════════════════════════════
        # FACTURATION EMPLOYÉ
        # PlatCategoriesalarie = tarif par catégorie salarié sur une période
        # Logique : pour chaque commande retirée :
        #   1. Trouver la catégorie salarié active de l'agent à la date du menu
        #      via UserCategoriesalarie (date_debut__date__lte, date_fin__date__gte)
        #   2. Trouver le PlatCategoriesalarie actif pour cette catégorie
        #      à la date du menu
        # ════════════════════════════════════════════════════
        total_montant_employe   = Decimal('0.00')
        fact_emp_par_mois       = {i: Decimal('0.00') for i in range(1, 13)}
        fact_emp_par_agence     = {}
        fact_emp_par_categorie  = {}

        for cmd in commandes_retirees:
            date_menu = cmd.menu.date_menu  # DateField → objet date Python

            # 1. Catégorie salarié active de l'agent à la date du menu
            user_cat = UserCategoriesalarie.objects.filter(
                user=cmd.user,
                date_debut__date__lte=date_menu,
            ).filter(
                Q(date_fin__isnull=True) | Q(date_fin__date__gte=date_menu)
            ).order_by('-date_debut').select_related('categoriesalarie').first()

            if not user_cat:
                continue

            # 2. PlatCategoriesalarie actif pour cette catégorie à la date du menu
            pc = PlatCategoriesalarie.objects.filter(
                categoriesalarie=user_cat.categoriesalarie,
                date_debut__date__lte=date_menu,
            ).filter(
                Q(date_fin__isnull=True) | Q(date_fin__date__gte=date_menu)
            ).order_by('-date_debut').first()

            if not pc:
                continue

            montant = pc.montant or Decimal('0.00')
            total_montant_employe += montant

            # Par mois
            mois = date_menu.month
            fact_emp_par_mois[mois] = fact_emp_par_mois.get(mois, Decimal('0.00')) + montant

            # Par agence
            agence_nom = cmd.menu.agence.nom_agence
            if agence_nom not in fact_emp_par_agence:
                fact_emp_par_agence[agence_nom] = {'agence': agence_nom, 'montant': Decimal('0.00'), 'nb_commandes': 0}
            fact_emp_par_agence[agence_nom]['montant']      += montant
            fact_emp_par_agence[agence_nom]['nb_commandes'] += 1

            # Par catégorie
            cat_lib = user_cat.categoriesalarie.libelle
            if cat_lib not in fact_emp_par_categorie:
                fact_emp_par_categorie[cat_lib] = {'categorie': cat_lib, 'montant': Decimal('0.00'), 'nb_commandes': 0}
            fact_emp_par_categorie[cat_lib]['montant']      += montant
            fact_emp_par_categorie[cat_lib]['nb_commandes'] += 1

        fact_employe = {
            'total_montant': float(total_montant_employe),
            'par_mois': {k: float(v) for k, v in fact_emp_par_mois.items()},
            'par_agence': [
                {'agence': v['agence'], 'montant': float(v['montant']), 'nb_commandes': v['nb_commandes']}
                for v in sorted(fact_emp_par_agence.values(), key=lambda x: x['montant'], reverse=True)
            ],
            'par_categorie': [
                {'categorie': v['categorie'], 'montant': float(v['montant']), 'nb_commandes': v['nb_commandes']}
                for v in sorted(fact_emp_par_categorie.values(), key=lambda x: x['montant'], reverse=True)
            ],
        }

        return Response({
            'kpis':             {
                'total_commandes':  total_commandes,
                'total_en_attente': total_en_attente,
                'total_retirees':   total_retirees,
                'total_annulees':   total_annulees,
                'total_menus':      total_menus,
                'taux_retrait':     taux_retrait,
                'taux_annulation':  taux_annulation,
            },
            'par_mois':         par_mois,
            'top_plats':        top_plats_data,
            'par_agence':       par_agence_data,
            'par_typeequipe':   par_typeequipe_data,
            'par_typeplat':     par_typeplat_data,
            'menus_recents':    menus_recents_data,
            'evolution_hebdo':  semaines,
            'evolution_par_date': evolution_par_date,
            'fact_fournisseur': fact_fournisseur,
            'fact_employe':     fact_employe,
        })


 
# ─────────────────────────────────────────────────────────────────────────────
# RETRAIT
# ─────────────────────────────────────────────────────────────────────────────

class RetraitViewSet(BaseViewSet):
    serializer_class = RetraitSerializer
    filterset_fields = ['menu', 'user']

    def get_queryset(self):
        return (
            Retrait.objects
            .select_related('user', 'menu__agence', 'menu__typeequipe')
            .order_by('-date_retrait')
        )


# ============================================================
# GÉNÉRATION FICHIER APB128
# ============================================================

class GenererAPB128View(APIView):
    parser_classes = [MultiPartParser]
    permission_classes = [AllowAny]

    def post(self, request):
        excel_file = request.FILES.get("file")
        if not excel_file:
            return Response({"error": "Fichier manquant."}, status=status.HTTP_400_BAD_REQUEST)

        try:
            wb       = load_workbook(filename=excel_file, data_only=True)
            ws_ident = wb["Identification"]
            ws_prel  = wb["Prelevements"]

            ident_data = {
                row[0].value: str(row[1].value).strip() if row[1].value is not None else ''
                for row in ws_ident.iter_rows(min_row=2)
            }

            prelevements = []
            for row in ws_prel.iter_rows(min_row=2, values_only=True):
                prelevements.append({
                    "Date_Execution": row[0], "Code_Guichet": row[1], "Compte": row[2],
                    "Cle_RIB": row[3], "Nom_Client": row[4], "Libelle_Banque_Client": row[5],
                    "Libelle_Operation": row[6], "Numero_Autorisation": row[7], "Montant": row[8]
                })

            def left_pad(val, size, char='0'):
                return str(val).strip()[:size].rjust(size, char)

            def right_pad(val, size, char=' '):
                return str(val).strip()[:size].ljust(size, char)

            def format_amount(val):
                val = int(val or 0)
                return str(val).rjust(12, '0')[-12:]

            lines         = []
            seq           = 1
            montant_total = 0
            compte_total  = int(ident_data["Compte"]) if ident_data["Compte"].isdigit() else 0

            lines.append(
                "03" + "1" + left_pad(seq, 5) +
                left_pad(ident_data["Date_Execution"], 6) + "CI006" +
                left_pad(ident_data["Code_Guichet"], 5) +
                left_pad(ident_data["Compte"], 12) +
                left_pad(ident_data["Cle_RIB"], 2) +
                right_pad(ident_data["Nom_Emetteur"], 24) +
                left_pad(ident_data["Code_Emetteur"], 5) +
                right_pad(" ", 1) +
                left_pad(ident_data["Date_Remise"], 6) +
                right_pad(ident_data["Reference_Remise"], 7) +
                right_pad(ident_data["Zone_Libre"], 47)
            )

            for p in prelevements:
                seq         += 1
                compte_val   = int(str(p["Compte"]).lstrip("0") or 0)
                montant_val  = int(p["Montant"] or 0)
                compte_total += compte_val
                montant_total += montant_val

                lines.append(
                    "03" + "2" + left_pad(seq, 5) +
                    left_pad(p["Date_Execution"], 6) + "CI006" +
                    left_pad(p["Code_Guichet"], 5) +
                    left_pad(p["Compte"], 12) +
                    left_pad(p["Cle_RIB"], 2) +
                    right_pad(p["Nom_Client"], 24) +
                    right_pad(p["Libelle_Banque_Client"], 17) +
                    right_pad(p["Libelle_Operation"], 20) +
                    left_pad(p["Numero_Autorisation"] or "0000000000", 10) +
                    format_amount(p["Montant"]) + "00" +
                    right_pad("", 5)
                )

            seq += 1
            lines.append(
                "03" + "9" + left_pad(seq, 5) +
                left_pad(ident_data["Date_Remise"], 6) +
                right_pad("", 8) +
                format_amount(compte_total) +
                right_pad(ident_data["Compte"].lstrip("0"), 77) +
                format_amount(montant_total) +
                right_pad("", 5)
            )

            output = BytesIO()
            output.write("\n".join(lines).encode('utf-8'))
            output.seek(0)
            return FileResponse(output, as_attachment=True, filename="FICHIER_APB128.txt")

        except Exception as e:
            return Response({"error": str(e)}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)


# ============================================================
# GESTION ENTRETIEN VÉHICULE
# ============================================================

class TypeVehiculeViewSet(BaseViewSet):
    queryset = TypeVehicule.objects.all().order_by('id')
    serializer_class = TypeVehiculeSerializer


class VehiculeViewSet(BaseViewSet):
    queryset = (
        Vehicule.objects
        .select_related('type_vehicule', 'agence', 'statut')
        .order_by('id')
    )
    serializer_class = VehiculeSerializer
    filterset_fields = ['immatriculation']


class PublicVehiculeListAPIView(ListAPIView):
    queryset = Vehicule.objects.all()
    serializer_class = VehiculeSerializer
    permission_classes = [AllowAny]
    pagination_class = None


class PublicTypeVehiculeListAPIView(ListAPIView):
    queryset = TypeVehicule.objects.all()
    serializer_class = TypeVehiculeSerializer
    permission_classes = [AllowAny]
    pagination_class = None


class ImportEntretienExcelView(APIView):
    parser_classes = [MultiPartParser]
    permission_classes = [AllowAny]

    def post(self, request):
        excel_file = request.FILES.get("file")
        if not excel_file:
            return Response({"error": "Fichier manquant"}, status=status.HTTP_400_BAD_REQUEST)

        try:
            wb    = load_workbook(excel_file, data_only=True)
            sheet = wb.active

            lignes_traitees = 0
            for row in sheet.iter_rows(min_row=2, values_only=True):
                immatricule_str, description, periode, debit, credit, solde = row

                if not immatricule_str or not periode:
                    continue

                try:
                    vehicule = Vehicule.objects.get(immatriculation=immatricule_str)
                except Vehicule.DoesNotExist:
                    continue

                EntretienVehicule.objects.filter(vehicule=vehicule, periode=periode).delete()
                EntretienVehicule.objects.create(
                    vehicule=vehicule,
                    description=description or '',
                    date_entretien_vehicule=now(),
                    periode=periode,
                    debit=Decimal(debit or 0),
                    credit=Decimal(credit or 0),
                    solde=Decimal(solde or 0)
                )
                lignes_traitees += 1

            return Response(
                {"success": f"{lignes_traitees} entretien(s) importé(s) avec succès"},
                status=status.HTTP_201_CREATED
            )

        except Exception as e:
            return Response({"error": str(e)}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)


class EntretienVehiculeViewSet(BaseViewSet):
    queryset = (
        EntretienVehicule.objects
        .select_related('vehicule', 'vehicule__agence')
        .order_by('periode')
    )
    serializer_class = EntretienVehiculeSerializer
    filterset_fields = ['vehicule']


class EtatEntretienGroupesView(APIView):
    permission_classes = [AllowAny]

    def get(self, request):
        annee           = request.query_params.get('annee')
        agence_id       = request.query_params.get('agence_id')
        immatriculation = request.query_params.get('immatriculation')

        if not annee:
            from datetime import datetime
            annee = datetime.today().year
        else:
            try:
                annee = int(annee)
            except ValueError:
                return Response({"error": "Année invalide"}, status=400)

        if agence_id:
            try:
                agence_id = int(agence_id)
            except ValueError:
                return Response({"error": "agence_id invalide"}, status=400)
        else:
            agence_id = None

        data = get_entretien_par_agence_avec_groupement(
            annee=annee, agence_id=agence_id, immatriculation=immatriculation
        )
        return Response(data)


# ============================================================
# RESSOURCES HUMAINES — Simulation salaire
# ============================================================

class SimulationSalaireUploadView(APIView):
    parser_classes = [MultiPartParser]
    permission_classes = [AllowAny]

    def post(self, request):
        from datetime import datetime

        fichier = request.FILES.get("file")
        if not fichier:
            return Response({"error": "Aucun fichier fourni"}, status=status.HTTP_400_BAD_REQUEST)

        try:
            df          = pd.read_excel(fichier)
            simulations = []

            for _, row in df.iterrows():
                try:
                    valeur_date = row["Date de naiss."]
                    if isinstance(valeur_date, (datetime, pd.Timestamp)):
                        date_naiss = valeur_date.date()
                    else:
                        date_naiss = datetime.strptime(str(valeur_date).strip(), "%d/%m/%Y").date()
                except Exception:
                    return Response({
                        "error": f"Date invalide pour le matricule {row['Matricule']}."
                    }, status=status.HTTP_400_BAD_REQUEST)

                age_2025 = 2025 - date_naiss.year
                salaire  = float(row["Salaire de base brut 2025"])

                simulation = SimulationSalaire.objects.create(
                    matricule=row["Matricule"],
                    categorie=row["Cat."],
                    date_naissance=date_naiss,
                    salaire_base=row["Salaire de base"],
                    sursalaire=row["Sursalaire"],
                    salaire_brut_2025=salaire
                )

                annee          = 2025
                salaire_actuel = salaire

                while age_2025 <= 60:
                    if salaire_actuel < 150000:
                        taux = 0.20
                    elif salaire_actuel < 250000:
                        taux = 0.10
                    else:
                        taux = 0.025

                    augmentation = round(salaire_actuel * taux)
                    cotisation   = round(salaire_actuel * 0.075) if salaire_actuel > 250000 else 0

                    SimulationHistorique.objects.create(
                        simulation=simulation, annee=annee, age=age_2025,
                        salaire=salaire_actuel, taux_augmentation=taux * 100,
                        augmentation=augmentation, cotisation=cotisation
                    )

                    salaire_actuel += augmentation
                    annee          += 1
                    age_2025       += 1

                simulations.append(SimulationSalaireSerializer(simulation).data)

            return Response(simulations, status=status.HTTP_201_CREATED)

        except Exception as e:
            return Response({"error": str(e)}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)


class SimulationSalaireListView(ListAPIView):
    queryset = SimulationSalaire.objects.all().order_by('-created_at')
    serializer_class = SimulationSalaireSerializer


class SimulationViewSet(BaseViewSet):
    queryset = SimulationSalaire.objects.all().order_by('matricule')
    serializer_class = SimulationSalaireSerializer
    permission_classes = [AllowAny]
    pagination_class = None

    def get_serializer_context(self):
        return {'request': self.request}

    def list(self, request, *args, **kwargs):
        from django.db.models import Q

        annee_debut = request.query_params.get('annee_debut')
        annee_fin   = request.query_params.get('annee_fin')
        salaire_min = request.query_params.get('salaire_min')
        salaire_max = request.query_params.get('salaire_max')

        filters = Q()
        if annee_debut: filters &= Q(annee__gte=annee_debut)
        if annee_fin:   filters &= Q(annee__lte=annee_fin)
        if salaire_min: filters &= Q(salaire__gte=salaire_min)
        if salaire_max: filters &= Q(salaire__lte=salaire_max)

        queryset = self.filter_queryset(self.get_queryset().prefetch_related('historique'))

        filtered_queryset = []
        for sim in queryset:
            historique_qs = sim.historique.all()
            if filters:
                historique_qs = historique_qs.filter(filters)
            if historique_qs.exists():
                filtered_queryset.append(sim)

        count = len(filtered_queryset)
        page  = self.paginate_queryset(filtered_queryset)
        if page is not None:
            serializer = self.get_serializer(page, many=True)
            data = serializer.data
        else:
            serializer = self.get_serializer(filtered_queryset, many=True)
            data = serializer.data

        historiques   = SimulationHistorique.objects.filter(filters) if filters else SimulationHistorique.objects.all()
        total_global  = {
            "salaire":      sum(float(h.salaire)      for h in historiques),
            "augmentation": sum(float(h.augmentation) for h in historiques),
            "cotisation":   sum(float(h.cotisation)   for h in historiques),
        }

        total_par_annee = {}
        for h in historiques:
            annee = str(h.annee)
            if annee not in total_par_annee:
                total_par_annee[annee] = {"salaire": 0, "augmentation": 0, "cotisation": 0, "effectif_aug": 0, "effectif_cot": 0}
            total_par_annee[annee]["salaire"]      += float(h.salaire)
            total_par_annee[annee]["augmentation"] += float(h.augmentation)
            total_par_annee[annee]["cotisation"]   += float(h.cotisation)
            if float(h.augmentation) > 0: total_par_annee[annee]["effectif_aug"] += 1
            if float(h.cotisation)   > 0: total_par_annee[annee]["effectif_cot"] += 1

        for annee, totals in total_par_annee.items():
            totals["effectif_non_aug"] = count - totals["effectif_aug"]
            totals["effectif_non_cot"] = count - totals["effectif_cot"]

        def init_tranche():
            return {"salaire": 0, "augmentation": 0, "cotisation": 0, "effectif_aug": 0, "effectif_cot": 0}

        total_par_annee_tranche_20  = {}
        total_par_annee_tranche_10  = {}
        total_par_annee_tranche_2_5 = {}

        for h in historiques:
            annee        = str(h.annee)
            salaire      = float(h.salaire)
            augmentation = float(h.augmentation)
            cotisation   = float(h.cotisation)

            if salaire < 150000:
                target = total_par_annee_tranche_20
            elif salaire <= 250000:
                target = total_par_annee_tranche_10
            else:
                target = total_par_annee_tranche_2_5

            if annee not in target:
                target[annee] = init_tranche()

            target[annee]["salaire"]      += salaire
            target[annee]["augmentation"] += augmentation
            target[annee]["cotisation"]   += cotisation
            if augmentation > 0: target[annee]["effectif_aug"] += 1
            if cotisation   > 0: target[annee]["effectif_cot"] += 1

        for total in [total_par_annee_tranche_20, total_par_annee_tranche_10, total_par_annee_tranche_2_5]:
            for annee, vals in total.items():
                vals["effectif_non_aug"] = count - vals["effectif_aug"]
                vals["effectif_non_cot"] = count - vals["effectif_cot"]

        return Response({
            "results": data, "total_global": total_global,
            "total_par_annee": total_par_annee,
            "total_par_annee_tranche_20": total_par_annee_tranche_20,
            "total_par_annee_tranche_10": total_par_annee_tranche_10,
            "total_par_annee_tranche_2_5": total_par_annee_tranche_2_5,
            "count": count
        })


class GenererSimulationHistorique2View(APIView):
    def post(self, request):
        deleted_count, _ = SimulationHistorique2.objects.all().delete()
        simulations  = SimulationSalaire.objects.all()
        total_crees  = 0

        for simulation in simulations:
            age_2025      = 2025 - simulation.date_naissance.year
            salaire       = Decimal(simulation.salaire_brut_2025).quantize(Decimal('1.'), rounding=ROUND_HALF_UP)
            annee         = 2025
            age           = age_2025
            is_first_year = True

            while age <= 60:
                if is_first_year and salaire < 150000:
                    augmentation = Decimal(150000 - salaire).quantize(Decimal('1.'), rounding=ROUND_HALF_UP)
                    taux         = (Decimal(augmentation) / salaire * 100).quantize(Decimal('0.01'))
                    salaire      = Decimal(150000)
                else:
                    if salaire < 150000:
                        taux = Decimal('20.00')
                    elif salaire <= 250000:
                        taux = Decimal('10.00')
                    else:
                        taux = Decimal('2.50')
                    augmentation = (salaire * taux / 100).quantize(Decimal('1.'), rounding=ROUND_HALF_UP)

                cotisation = (salaire * Decimal('0.075')).quantize(Decimal('1.')) if salaire > 250000 else Decimal(0)

                SimulationHistorique2.objects.create(
                    simulation=simulation, annee=annee, age=age, salaire=salaire,
                    taux_augmentation=taux, augmentation=augmentation, cotisation=cotisation
                )

                salaire      += augmentation
                annee        += 1
                age          += 1
                total_crees  += 1
                is_first_year = False

        return Response({
            "message": f"{deleted_count} ancienne(s) ligne(s) supprimée(s). "
                       f"{total_crees} nouvelle(s) ligne(s) créée(s) dans SimulationHistorique2."
        }, status=status.HTTP_201_CREATED)


class Simulation2ViewSet(BaseViewSet):
    queryset = SimulationSalaire.objects.all().order_by('matricule')
    serializer_class = SimulationSalaire2Serializer
    permission_classes = [AllowAny]
    pagination_class = None

    def get_serializer_context(self):
        return {'request': self.request}

    def list(self, request, *args, **kwargs):
        from django.db.models import Q

        annee_debut = request.query_params.get('annee_debut')
        annee_fin   = request.query_params.get('annee_fin')
        salaire_min = request.query_params.get('salaire_min')
        salaire_max = request.query_params.get('salaire_max')

        filters = Q()
        if annee_debut: filters &= Q(annee__gte=annee_debut)
        if annee_fin:   filters &= Q(annee__lte=annee_fin)
        if salaire_min: filters &= Q(salaire__gte=salaire_min)
        if salaire_max: filters &= Q(salaire__lte=salaire_max)

        queryset = self.filter_queryset(self.get_queryset().prefetch_related('historique2'))

        filtered_queryset = []
        for sim in queryset:
            historique_qs = sim.historique2.all()
            if filters:
                historique_qs = historique_qs.filter(filters)
            if historique_qs.exists():
                filtered_queryset.append(sim)

        count = len(filtered_queryset)
        page  = self.paginate_queryset(filtered_queryset)
        if page is not None:
            serializer = self.get_serializer(page, many=True)
            data = serializer.data
        else:
            serializer = self.get_serializer(filtered_queryset, many=True)
            data = serializer.data

        historiques  = SimulationHistorique2.objects.filter(filters) if filters else SimulationHistorique2.objects.all()
        total_global = {
            "salaire":      sum(float(h.salaire)      for h in historiques),
            "augmentation": sum(float(h.augmentation) for h in historiques),
            "cotisation":   sum(float(h.cotisation)   for h in historiques),
        }

        total_par_annee = {}
        for h in historiques:
            annee = str(h.annee)
            if annee not in total_par_annee:
                total_par_annee[annee] = {"salaire": 0, "augmentation": 0, "cotisation": 0, "effectif_aug": 0, "effectif_cot": 0}
            total_par_annee[annee]["salaire"]      += float(h.salaire)
            total_par_annee[annee]["augmentation"] += float(h.augmentation)
            total_par_annee[annee]["cotisation"]   += float(h.cotisation)
            if float(h.augmentation) > 0: total_par_annee[annee]["effectif_aug"] += 1
            if float(h.cotisation)   > 0: total_par_annee[annee]["effectif_cot"] += 1

        for annee, totals in total_par_annee.items():
            totals["effectif_non_aug"] = count - totals["effectif_aug"]
            totals["effectif_non_cot"] = count - totals["effectif_cot"]

        def init_tranche():
            return {"salaire": 0, "augmentation": 0, "cotisation": 0, "effectif_aug": 0, "effectif_cot": 0}

        total_par_annee_tranche_20  = {}
        total_par_annee_tranche_10  = {}
        total_par_annee_tranche_2_5 = {}

        for h in historiques:
            annee        = str(h.annee)
            salaire      = float(h.salaire)
            augmentation = float(h.augmentation)
            cotisation   = float(h.cotisation)

            if salaire < 150000:
                target = total_par_annee_tranche_20
            elif salaire <= 250000:
                target = total_par_annee_tranche_10
            else:
                target = total_par_annee_tranche_2_5

            if annee not in target:
                target[annee] = init_tranche()

            target[annee]["salaire"]      += salaire
            target[annee]["augmentation"] += augmentation
            target[annee]["cotisation"]   += cotisation
            if augmentation > 0: target[annee]["effectif_aug"] += 1
            if cotisation   > 0: target[annee]["effectif_cot"] += 1

        for total in [total_par_annee_tranche_20, total_par_annee_tranche_10, total_par_annee_tranche_2_5]:
            for annee, vals in total.items():
                vals["effectif_non_aug"] = count - vals["effectif_aug"]
                vals["effectif_non_cot"] = count - vals["effectif_cot"]

        return Response({
            "results": data, "total_global": total_global,
            "total_par_annee": total_par_annee,
            "total_par_annee_tranche_20": total_par_annee_tranche_20,
            "total_par_annee_tranche_10": total_par_annee_tranche_10,
            "total_par_annee_tranche_2_5": total_par_annee_tranche_2_5,
            "count": count
        })


# ============================================================
# PLAN COMPTABLE
# ============================================================

class CategorieComptableViewSet(BaseViewSet):
    queryset = CategorieComptable.objects.all().order_by('id')
    serializer_class = CategorieComptableSerializer


class PublicCategorieComptableListAPIView(ListAPIView):
    queryset = CategorieComptable.objects.all()
    serializer_class = CategorieComptableSerializer
    permission_classes = [AllowAny]
    pagination_class = None


class ClasseComptableViewSet(BaseViewSet):
    queryset = ClasseComptable.objects.all().order_by('id')
    serializer_class = ClasseComptableSerializer
    filterset_fields = ['libelle']


class ClasseComptableListAPIView(ListAPIView):
    queryset = ClasseComptable.objects.all()
    serializer_class = ClasseComptableSerializer
    permission_classes = [AllowAny]
    pagination_class = None


class PosteReportingViewSet(BaseViewSet):
    queryset = PosteReporting.objects.all().order_by('id')
    serializer_class = PosteReportingSerializer
    filterset_fields = ['libelle']


class PosteReportingListAPIView(ListAPIView):
    queryset = PosteReporting.objects.all()
    serializer_class = PosteReportingSerializer
    permission_classes = [AllowAny]
    pagination_class = None


class CompteComptableViewSet(BaseViewSet):
    queryset = (
        CompteComptable.objects
        .select_related('classecomptable', 'postereporting')
        .order_by('id')
    )
    serializer_class = CompteComptableSerializer
    filterset_fields = ['libelle']


class CompteComptableListAPIView(ListAPIView):
    queryset = CompteComptable.objects.all()
    serializer_class = CompteComptableSerializer
    permission_classes = [AllowAny]
    pagination_class = None


class FonctionListAPIView(ListAPIView):
    queryset = Fonction.objects.all()
    serializer_class = FonctionSerializer
    permission_classes = [AllowAny]
    pagination_class = None


class ServiceListAPIView(ListAPIView):
    queryset = Service.objects.all()
    serializer_class = ServiceSerializer
    permission_classes = [AllowAny]
    pagination_class = None


class PosteListAPIView(ListAPIView):
    queryset = Poste.objects.select_related('fonction', 'service').all()
    serializer_class = PosteSerializer
    permission_classes = [AllowAny]
    pagination_class = None


class CategoriesalarieListAPIView(ListAPIView):
    queryset = Categoriesalarie.objects.all()
    serializer_class = CategoriesalarieSerializer
    permission_classes = [AllowAny]
    pagination_class = None


class GroupListAPIView(ListAPIView):
    queryset = Group.objects.all()
    serializer_class = GroupSerializer
    permission_classes = [AllowAny]
    pagination_class = None

class AllPlatNopginListAPIView(ListAPIView):
    queryset = Plat.objects.all()
    serializer_class = PlatSerializer
    permission_classes = [IsAuthenticated]
    pagination_class   = None


# ============================================================
# AJOUTER DANS views.py
# ============================================================

# ── Imports à ajouter en haut de views.py ───────────────────
# from .models import DeviceToken, Notification
# from .fcm_service import envoyer_notification_user

# ══════════════════════════════════════════════════════════════
# DEVICE TOKEN — Enregistrement des tokens FCM
# ══════════════════════════════════════════════════════════════

class DeviceTokenView(APIView):
    """
    POST /api/api/device-tokens/
    Enregistre ou met à jour le token FCM de l'appareil connecté.
    Appelé automatiquement par Flutter après login.
    """
    permission_classes = [IsAuthenticated]

    def post(self, request):
        token    = request.data.get('token', '').strip()
        platform = request.data.get('platform', 'android')

        if not token:
            return Response(
                {'error': 'Token requis.'},
                status=status.HTTP_400_BAD_REQUEST
            )

        if platform not in ['android', 'ios', 'web']:
            return Response(
                {'error': 'Plateforme invalide. Valeurs : android, ios, web.'},
                status=status.HTTP_400_BAD_REQUEST
            )

        # update_or_create : évite les doublons si le token existe déjà
        # pour un autre user (token recyclé par FCM)
        DeviceToken.objects.filter(token=token).exclude(
            user=request.user
        ).delete()  # supprimer si token appartient à un autre user

        obj, created = DeviceToken.objects.update_or_create(
            token=token,
            defaults={
                'user':     request.user,
                'platform': platform,
            }
        )

        return Response(
            {
                'message':  'Token enregistré.' if created else 'Token mis à jour.',
                'id':       obj.id,
                'platform': obj.platform,
            },
            status=status.HTTP_201_CREATED if created else status.HTTP_200_OK
        )

    def delete(self, request):
        """
        DELETE /api/api/device-tokens/supprimer/?token=xxx
        Supprime le token FCM au logout.
        """
        token = request.query_params.get('token', '').strip()
        if not token:
            return Response(
                {'error': 'Token requis.'},
                status=status.HTTP_400_BAD_REQUEST
            )
        deleted, _ = DeviceToken.objects.filter(
            user=request.user, token=token
        ).delete()
        return Response(
            {'message': f'{deleted} token(s) supprimé(s).'},
            status=status.HTTP_200_OK
        )


# ══════════════════════════════════════════════════════════════
# NOTIFICATIONS — CRUD
# ══════════════════════════════════════════════════════════════

class NotificationListView(APIView):
    """
    GET /api/api/notifications/mes-notifications/
    Retourne les notifications de l'user connecté
    avec le compteur de non lues.
    """
    permission_classes = [IsAuthenticated]

    def get(self, request):
        notifs = Notification.objects.filter(
            user=request.user
        ).order_by('-created_at')[:50]  # 50 dernières max

        non_lues = Notification.objects.filter(
            user=request.user, lu=False
        ).count()

        data = [
            {
                'id':           n.id,
                'type':         n.type,
                'titre':        n.titre,
                'message':      n.message,
                'lu':           n.lu,
                'commande_id':  n.commande_id,
                'menu_id':      n.menu_id,
                'created_at':   n.created_at.isoformat(),
            }
            for n in notifs
        ]

        return Response({
            'non_lues':      non_lues,
            'notifications': data,
        })


class NotificationDetailView(APIView):
    """
    PATCH /api/api/notifications/{id}/marquer-lu/
    DELETE /api/api/notifications/{id}/supprimer/
    """
    permission_classes = [IsAuthenticated]

    def _get_notif(self, request, pk):
        try:
            return Notification.objects.get(pk=pk, user=request.user)
        except Notification.DoesNotExist:
            return None

    def patch(self, request, pk):
        notif = self._get_notif(request, pk)
        if not notif:
            return Response(status=status.HTTP_404_NOT_FOUND)
        notif.lu = True
        notif.save(update_fields=['lu'])
        return Response({'message': 'Notification marquée comme lue.'})

    def delete(self, request, pk):
        notif = self._get_notif(request, pk)
        if not notif:
            return Response(status=status.HTTP_404_NOT_FOUND)
        notif.delete()
        return Response({'message': 'Notification supprimée.'})


class NotificationBulkView(APIView):
    """
    PATCH /api/api/notifications/tout-marquer-lu/
    DELETE /api/api/notifications/tout-supprimer/
    """
    permission_classes = [IsAuthenticated]

    def patch(self, request):
        count = Notification.objects.filter(
            user=request.user, lu=False
        ).update(lu=True)
        return Response({'message': f'{count} notification(s) marquée(s) comme lues.'})

    def delete(self, request):
        count, _ = Notification.objects.filter(user=request.user).delete()
        return Response({'message': f'{count} notification(s) supprimée(s).'})
